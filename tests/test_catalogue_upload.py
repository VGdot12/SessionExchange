import tempfile
import unittest
import io
from contextlib import redirect_stderr, redirect_stdout
from pathlib import Path
from datetime import date, time

from openpyxl import Workbook
from session_swap.catalogue_upload import read_uploaded_workbook, main
from session_swap.service import Problem

SUBJECT = 'Upcoming Event  | PGP Co 2027 | Pre Placement Talk : Session Scheduled'
HEADERS = ['Subject\u00a0', 'Company Name', 'Industry\u00a0', 'CBD Name', 'Event Type', 'Event Date', 'Event Time', 'Duration', 'Mode', 'Venue', 'Received On']


class CatalogueUploadTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.path = Path(self.temp.name) / 'sessions.xlsx'

    def tearDown(self):
        self.temp.cleanup()

    def write(self, headers, rows, name='Table1'):
        wb = Workbook()
        wb.active.title = name
        wb.active.append(headers)
        for row in rows:
            wb.active.append(row)
        wb.save(self.path)
        wb.close()

    def invite(self, **changes):
        values = dict(zip(HEADERS, [SUBJECT, 'Nagarro', 'PRIVATE INDUSTRY', 'PRIVATE PERSON', 'Pre Placement Talk', '27 October, 2026', '21:30 - 22:30', '1 hr', 'Offline', 'LT 3', 'PRIVATE RECEIVED']))
        values.update(changes)
        return [values[h] for h in HEADERS]

    def test_actual_headings_named_date_and_private_columns_are_handled(self):
        self.write(HEADERS, [self.invite()])
        payload = read_uploaded_workbook(self.path)
        event = payload['sessions'][0]
        self.assertEqual(event['title'], 'Nagarro · Pre Placement Talk')
        self.assertEqual((event['room'], event['date'], event['start_time'], event['end_time']), ('LT 3', '2026-10-27', '21:30', '22:30'))
        self.assertNotIn('PRIVATE', str(payload))
        self.assertNotIn('Subject', str(payload))

    def test_excel_date_and_simplified_format_keep_stable_ids(self):
        self.write(HEADERS, [self.invite(**{'Event Date': date(2026, 10, 27)})])
        first = read_uploaded_workbook(self.path)['sessions'][0]
        self.write(['title', 'room', 'time', 'date'], [['Nagarro · Pre Placement Talk', 'LT 4', '21:30 – 23:00', '27/10/2026']], 'Sessions')
        second = read_uploaded_workbook(self.path)['sessions'][0]
        self.assertEqual(first['session_id'], second['session_id'])
        self.assertEqual(second['room'], 'LT 4')

    def test_all_valid_rows_import_regardless_of_subject_and_identical_rows_collapse(self):
        self.write(HEADERS, [self.invite(), self.invite(), self.invite(**{'Subject\u00a0': SUBJECT.replace('2027', '2028'), 'Company Name': 'Other'})])
        self.assertEqual(len(read_uploaded_workbook(self.path)['sessions']), 2)
        self.write(HEADERS, [self.invite(**{'Subject\u00a0': 'A1'})])
        self.assertEqual(len(read_uploaded_workbook(self.path)['sessions']), 1)

    def test_duration_and_subject_are_unused_even_if_blank_or_formula(self):
        for duration in ['1hr', time(10, 30), None, '=PRIVATE_DURATION()']:
            self.write(HEADERS, [self.invite(**{'Subject\u00a0': '=PRIVATE_SUBJECT()', 'Duration': duration, 'Event Time': '8:30 - 9:30'})])
            event = read_uploaded_workbook(self.path)['sessions'][0]
            self.assertEqual((event['start_time'], event['end_time']), ('08:30', '09:30'))
            self.assertNotIn('PRIVATE', str(event))

    def test_supplied_layout_with_sample_subject_and_correct_time_range_imports(self):
        self.write([''], [[''], [None, *HEADERS], [None, *self.invite(**{'Subject\u00a0': 'A1', 'Event Time': '8:30 - 9:30', 'Duration': '1hr'})]], 'SessionExchange')
        event = read_uploaded_workbook(self.path)['sessions'][0]
        self.assertEqual((event['start_time'], event['end_time']), ('08:30', '09:30'))

    def test_single_excel_start_time_requires_a_range_and_never_uses_duration_as_end(self):
        self.write([''], [[''], [None, *HEADERS], [None, *self.invite(**{'Subject\u00a0': 'A1', 'Event Time': time(8, 30), 'Duration': time(10, 30)})]], 'SessionExchange')
        status, error, target = self.run_reader()
        self.assertEqual(status, 1)
        self.assertIn('Row 4', error)
        self.assertIn('08:30 - 09:30', error)
        self.assertIn('Duration is not used', error)
        self.assertFalse(target.exists())

    def test_invalid_rows_formulas_and_conflicting_duplicate_events_fail(self):
        for changes in [{'Event Time': '23:00 - 01:00'}, {'Event Date': '2026-02-30'}, {'Venue': '=A1'}, {'Company Name': ''}]:
            self.write(HEADERS, [self.invite(**changes)])
            with self.assertRaises(Problem):
                read_uploaded_workbook(self.path)
        self.write(HEADERS, [self.invite(), self.invite(**{'Venue': 'Other room'})])
        with self.assertRaises(Problem):
            read_uploaded_workbook(self.path)

    def test_original_schema_and_optional_stable_ids_are_supported(self):
        from session_swap.importer import HEADERS as ORIGINAL
        self.write(ORIGINAL['Sessions'], [['legacy', 'topic', 'Legacy', 'A', '2026-10-27', '20:00', '21:00', 'LT 1', 'PRIVATE PERSON']])
        event = read_uploaded_workbook(self.path)['sessions'][0]
        self.assertEqual((event['session_id'], event['course_id'], event['instructor']), ('legacy', 'topic', ''))
        self.write(['title', 'room', 'time', 'date', 'session_id'], [['Stable event', 'LT 1', '20:00 - 21:00', '2026-10-27', 'stable-1']])
        self.assertEqual(read_uploaded_workbook(self.path)['sessions'][0]['session_id'], 'stable-1')

    def test_ambiguous_sheets_fail_and_named_sessions_sheet_wins(self):
        self.write(['title', 'room', 'time', 'date'], [['One', 'LT 1', '20:00 - 21:00', '2026-10-27']])
        from openpyxl import load_workbook
        wb = load_workbook(self.path)
        wb.copy_worksheet(wb.active)
        wb.save(self.path)
        wb.close()
        with self.assertRaises(Problem):
            read_uploaded_workbook(self.path)
        wb = load_workbook(self.path)
        wb.active.title = 'Sessions'
        wb.save(self.path)
        wb.close()
        self.assertEqual(len(read_uploaded_workbook(self.path)['sessions']), 1)

    def run_reader(self):
        err, out = io.StringIO(), io.StringIO()
        target = self.path.with_suffix('.json')
        with redirect_stderr(err), redirect_stdout(out):
            status = main([str(self.path), str(target)])
        return status, err.getvalue(), target

    def test_workflow_errors_identify_date_and_row_without_logging_cell_contents(self):
        self.write(HEADERS, [self.invite(**{'Event Date': 'PRIVATE INVALID DATE'})])
        status, error, target = self.run_reader()
        self.assertEqual(status, 1)
        self.assertIn('Row 2', error)
        self.assertIn('Event Date', error)
        self.assertNotIn('PRIVATE', error)
        self.assertFalse(target.exists())

    def test_header_row_is_detected_below_blank_rows_and_after_an_empty_first_column(self):
        self.write([''], [[''], [None, *HEADERS], [None, *self.invite()]], 'SessionExchange')
        event = read_uploaded_workbook(self.path)['sessions'][0]
        self.assertEqual(event['date'], '2026-10-27')
        self.assertEqual(event['room'], 'LT 3')

    def test_detected_header_row_keeps_actual_excel_row_numbers_in_errors(self):
        self.write(['Session catalogue'], [[], [*HEADERS], self.invite(**{'Event Date': 'PRIVATE INVALID DATE'})])
        status, error, _ = self.run_reader()
        self.assertEqual(status, 1)
        self.assertIn('Row 4', error)
        self.assertNotIn('PRIVATE', error)

    def test_workflow_identifies_missing_headers_and_does_not_filter_other_subjects(self):
        self.write(['Company Name', 'Event Date', 'Event Time', 'Venue'], [['PRIVATE', '2026-10-27', '20:00 - 21:00', 'LT 1']])
        status, error, _ = self.run_reader()
        self.assertEqual(status, 1)
        self.assertIn('Missing invite headings: event type', error)
        self.assertNotIn('PRIVATE', error)
        self.write(HEADERS, [self.invite(**{'Subject\u00a0': 'PRIVATE OTHER EMAIL'})])
        status, error, _ = self.run_reader()
        self.assertEqual(status, 0)
        self.assertEqual(error, '')
        self.assertNotIn('PRIVATE', error)

    def test_workflow_identifies_conflicting_row_numbers_and_formula_column(self):
        self.write(HEADERS, [self.invite(), self.invite(**{'Venue': 'PRIVATE ROOM'})])
        status, error, _ = self.run_reader()
        self.assertEqual(status, 1)
        self.assertIn('Rows 2 and 3', error)
        self.assertNotIn('PRIVATE', error)
        self.write(HEADERS, [self.invite(**{'Venue': '=PRIVATE_FORMULA()'})])
        _, error, _ = self.run_reader()
        self.assertIn('Row 2, venue', error)
        self.assertNotIn('PRIVATE', error)

    def test_workflow_reports_wrong_file_format_and_success_writes_only_event_fields(self):
        self.path.write_text('PRIVATE NOT AN EXCEL FILE')
        status, error, _ = self.run_reader()
        self.assertEqual(status, 1)
        self.assertIn('not a readable .xlsx workbook', error)
        self.assertNotIn('PRIVATE', error)
        self.write(HEADERS, [self.invite()])
        status, error, target = self.run_reader()
        self.assertEqual(status, 0)
        self.assertEqual(error, '')
        self.assertNotIn('PRIVATE', target.read_text())


if __name__ == '__main__':
    unittest.main()
